#!/usr/bin/env python3
"""Build finished PDF, Word, text, and print assets from the canonical Markdown files."""
from __future__ import annotations
from pathlib import Path
import re, textwrap, xml.sax.saxutils
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether, PageBreak, HRFlowable
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parent.parent
PURPLE = colors.HexColor("#7753CE")
INK = colors.HexColor("#302D3B")
MUTED = colors.HexColor("#716B7D")
PAPER = colors.HexColor("#F8F7F4")
LILAC = colors.HexColor("#EEE9F6")
BLUE = colors.HexColor("#3D7D9F")
DOCS = list((ROOT/"01_Google_Drive").rglob("*.md"))

styles = getSampleStyleSheet()
styles.add(ParagraphStyle(name="OrgTitle", fontName="Helvetica-Bold", fontSize=23, leading=29, textColor=INK, spaceAfter=12))
styles.add(ParagraphStyle(name="OrgH2", fontName="Helvetica-Bold", fontSize=12, leading=16, textColor=PURPLE, spaceBefore=16, spaceAfter=7, keepWithNext=True))
styles.add(ParagraphStyle(name="OrgBody", fontName="Helvetica", fontSize=10.2, leading=15.5, textColor=INK, spaceAfter=8))
styles.add(ParagraphStyle(name="OrgNote", fontName="Helvetica", fontSize=8.7, leading=12.5, textColor=MUTED, spaceAfter=5))
styles.add(ParagraphStyle(name="OrgList", parent=styles["OrgBody"], leftIndent=20, firstLineIndent=-12, spaceAfter=5))

def safe(s):
    return xml.sax.saxutils.escape(s, {"\"":"&quot;"})

def pars(lines):
    buf=[]
    for line in lines:
        if line.strip():
            buf.append(line.strip())
        elif buf:
            yield " ".join(buf); buf=[]
    if buf:
        yield " ".join(buf)

def markdown_lines(raw):
    # Keep document headings and prose exactly as authored; only convert markup.
    for block in re.split(r"\n\s*\n", raw.strip()):
        if not block.strip(): continue
        lines=[x.strip() for x in block.splitlines() if x.strip()]
        if not lines: continue
        for line in lines:
            if line.startswith("# "):
                yield ("title", line[2:])
            elif line.startswith("## "):
                yield ("heading", line[3:])
            elif re.match(r"^\d+\. ",line):
                yield ("list",line)
            elif line.startswith("- ["):
                yield ("list",line.replace("- [ ]", "□").replace("- [x]", "☑"))
            elif line.startswith("- "):
                yield ("list","• "+line[2:])
            elif len(lines)>1 and any(z.startswith("## ") for z in lines):
                yield ("body",line)
            elif len(lines)>1 and not lines[0].startswith("#"):
                yield ("body",line)
            else:
                yield ("body",line)

def draw_pdf_footer(canvas, doc):
    canvas.saveState()
    width,height = A4
    canvas.setFillColor(PURPLE)
    canvas.rect(0,height-9,width,9,stroke=0,fill=1)
    canvas.setStrokeColor(colors.HexColor("#DED9E6"))
    canvas.line(48,45,width-48,45)
    canvas.setFont("Helvetica-Bold",7.4)
    canvas.setFillColor(MUTED)
    canvas.drawString(48,31,"CENTRE VALLEY COMMUNITY PANTRY  /  ORGFLOW DEMONSTRATION")
    canvas.setFont("Helvetica",7.2)
    canvas.drawRightString(width-48,31,f"FICTIONAL TRAINING CONTENT  |  {doc.page}")
    canvas.restoreState()

def render_pdf(md_path):
    text=md_path.read_text()
    target=md_path.with_suffix(".pdf")
    doc=SimpleDocTemplate(str(target),pagesize=A4,leftMargin=53,rightMargin=53,
                          topMargin=60,bottomMargin=60,title=md_path.stem.replace("_"," "),
                          author="OrgFlow demonstration")
    flow=[]
    for kind,value in markdown_lines(text):
        v=safe(value)
        if kind=="title":
            flow.append(Paragraph(v,styles["OrgTitle"]))
            flow.append(HRFlowable(width="100%",thickness=1.2,color=LILAC,spaceAfter=10))
        elif kind=="heading":
            flow.append(Paragraph(v,styles["OrgH2"]))
        elif kind=="list":
            v=v.replace("□","[ ]").replace("☑","[x]").replace("•","-")
            flow.append(Paragraph(v,styles["OrgList"]))
        else:
            st=styles["OrgNote"] if (value.startswith(("Document owner:","Owner:","Approved:","Prepared","Review due:","Next review:","Transition date:","Last updated:","Status:","Reporting period:","Event date:","Meeting example:","Successor:","Venue:","Leaving date:","Fictional"))) else styles["OrgBody"]
            flow.append(Paragraph(v,st))
    doc.build(flow,onFirstPage=draw_pdf_footer,onLaterPages=draw_pdf_footer)
    md_path.with_suffix(".txt").write_text(text, encoding="utf-8")
    return target

def make_docx(path):
    content=path.read_text(encoding="utf-8")
    doc=Document()
    section=doc.sections[0]
    section.top_margin=Inches(.83); section.bottom_margin=Inches(.72)
    section.left_margin=Inches(.86); section.right_margin=Inches(.86)
    normal=doc.styles["Normal"]; normal.font.name="Aptos"; normal.font.size=Pt(10.5)
    normal.font.color.rgb=RGBColor(48,45,59)
    normal.paragraph_format.space_after=Pt(8)
    for kind,line in markdown_lines(content):
        if kind=="title":
            p=doc.add_heading(line,0)
            for r in p.runs:r.font.color.rgb=RGBColor(48,45,59)
        elif kind=="heading":
            p=doc.add_heading(line,2)
            for r in p.runs:r.font.color.rgb=RGBColor(119,83,206)
        elif kind=="list":
            p=doc.add_paragraph(line,style="List Number" if re.match(r"^\d+\. ",line) else "List Bullet")
        else:
            doc.add_paragraph(line)
    footer=section.footer.paragraphs[0]
    footer.text="CENTRE VALLEY COMMUNITY PANTRY  |  FICTIONAL ORGFLOW DEMONSTRATION"
    footer.style="Caption"
    out=path.with_suffix(".docx")
    doc.save(out)
    return out

def draw_poster(path):
    w,h=A4;c=Canvas(str(path),pagesize=A4)
    c.setFillColor(PAPER);c.rect(0,0,w,h,fill=1,stroke=0)
    c.setFillColor(PURPLE);c.roundRect(45,h-95,w-90,6,3,fill=1,stroke=0)
    c.setFont("Helvetica-Bold",14);c.drawString(52,h-140,"CENTRE VALLEY COMMUNITY PANTRY")
    c.setFillColor(INK);c.setFont("Helvetica-Bold",53)
    c.drawString(52,h-270,"Harvest")
    c.drawString(52,h-334,"Dinner")
    c.setFillColor(PURPLE);c.setFont("Helvetica",18)
    c.drawString(55,h-389,"A gathering for our community")
    c.setFillColor(LILAC);c.roundRect(50,h-592,w-100,155,22,fill=1,stroke=0)
    c.setFillColor(INK);c.setFont("Helvetica-Bold",23)
    c.drawString(74,h-491,"Saturday, October 24, 2026")
    c.setFont("Helvetica",18);c.drawString(74,h-524,"Doors at 5:30 PM  |  Program at 6:15 PM")
    c.setFont("Helvetica",16);c.drawString(74,h-553,"Fictional community hall, State College")
    c.setFont("Helvetica-Bold",19);c.drawString(54,h-659,"An evening to connect, listen and say thank you.")
    c.setFont("Helvetica",13);c.setFillColor(MUTED)
    c.drawString(54,90,"Invitation concept created for an OrgFlow software demonstration.")
    c.drawString(54,70,"No real event, sale, registration, or donation is being offered.")
    c.setFillColor(PURPLE);c.rect(0,0,w,17,fill=1,stroke=0)
    c.setTitle("Harvest Dinner - fictional event invitation")
    c.showPage();c.save()

def draw_report(path):
    w,h=A4;c=Canvas(str(path),pagesize=A4)
    c.setFillColor(PAPER);c.rect(0,0,w,h,fill=1,stroke=0)
    c.setFillColor(PURPLE);c.rect(0,h-15,w,15,fill=1,stroke=0)
    c.setFillColor(INK);c.setFont("Helvetica-Bold",28);c.drawString(50,h-86,"Quarterly program snapshot")
    c.setFont("Helvetica",14);c.setFillColor(MUTED);c.drawString(50,h-115,"July to September 2026  |  Fictional demonstration")
    c.setFillColor(PURPLE);c.setFont("Helvetica-Bold",72);c.drawString(52,h-235,"1,260")
    c.setFont("Helvetica",17);c.setFillColor(INK);c.drawString(57,h-263,"example household visits")
    c.setStrokeColor(colors.HexColor("#DEDAE7"));c.line(52,h-307,w-52,h-307)
    values=[401,419,440];names=["July","August","September"]
    maxv=480;origin=410;baseline=233
    for i,(name,value) in enumerate(zip(names,values)):
      x=66+i*165
      c.setFillColor(LILAC);c.roundRect(x,baseline,110,210,12,fill=1,stroke=0)
      bar_h=value/maxv*200
      c.setFillColor(PURPLE);c.roundRect(x,baseline,110,bar_h,12,fill=1,stroke=0)
      c.setFillColor(INK);c.setFont("Helvetica-Bold",19);c.drawCentredString(x+55,baseline-31,str(value))
      c.setFont("Helvetica",12);c.setFillColor(MUTED);c.drawCentredString(x+55,baseline-50,name)
    c.setFont("Helvetica",11);c.setFillColor(MUTED)
    for i,line in enumerate(["These figures are invented solely for software evaluation and filming.",
       "A household visit is a visit, not a count of distinct households.",
       "For a real grant report, reconcile data and obtain human approval."]):
      c.drawString(53,135-i*22,line)
    c.setFillColor(PURPLE);c.rect(0,0,w,15,fill=1,stroke=0)
    c.showPage();c.save()

if __name__=="__main__":
    pdfs=[render_pdf(x) for x in DOCS]
    # Word counterparts are genuine editable documents for varied Drive display.
    choices=["Volunteer_Welcome_and_First_Shift","Donor_Thank_You_and_Receipts","Maya_Chen_Transfer_Guide","Quarterly_Grant_Reporting"]
    word=[make_docx(x) for x in DOCS if x.stem in choices]
    draw_poster(ROOT/"01_Google_Drive/Programs/Harvest_Dinner_Invitation.pdf")
    draw_report(ROOT/"01_Google_Drive/Programs/Quarterly_Impact_Visual.pdf")
    print("SOURCE_DOCUMENTS",len(DOCS),"PDFS",len(pdfs)+2,"WORD_DOCUMENTS",len(word),"TEXT_FILES",len(DOCS))
    for p in pdfs: print("PDF",p.relative_to(ROOT),round(p.stat().st_size/1024,1),"KB")
